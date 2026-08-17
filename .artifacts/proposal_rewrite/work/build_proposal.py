from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/home/wielio/Documents/Wiefran/Proyek/REKSA/REKSA")
OUT = ROOT / "output" / "documents" / "REKSA_Proposal_GEMASTIK_XIX_2026_Revisi_Monokrom.docx"
MEDIA = ROOT / ".artifacts" / "proposal_rewrite" / "work" / "media" / "word" / "media"

NAVY = "000000"
BLUE = "000000"
LIGHT_BLUE = "FFFFFF"
PALE = "FFFFFF"
GOLD = "000000"
GRAY = "000000"
MID_GRAY = "B7B7B7"
WHITE = "FFFFFF"
BLACK = "000000"
RED = "000000"
GREEN = "000000"
TABLE_RULE = "808080"


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    # Gaya proposal akademis monokrom: seluruh teks, termasuk caption dan
    # metadata, dicetak hitam. Parameter color dipertahankan agar pemanggil
    # lama tetap kompatibel, tetapi tidak dipakai sebagai aksen warna.
    run.font.color.rgb = RGBColor.from_string(BLACK)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size="6", style="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), style)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_twips = round(sum(widths_cm) / 2.54 * 1440)
    column_twips = [round(width / 2.54 * 1440) for width in widths_cm]
    column_twips[-1] += total_twips - sum(column_twips)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for twips in column_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(twips))
        grid.append(col)
    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths_cm)):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(column_twips[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_picture_with_alt(run, path, width, alt_text):
    shape = run.add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return shape


def add_table(doc, headers, rows, widths, font_size=9.2, header_size=9.2, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_widths(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade_cell(cell, WHITE)
        set_cell_border(cell, BLACK, "8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(text))
        set_run_font(r, size=header_size, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, value in enumerate(values):
            cell = cells[col_idx]
            shade_cell(cell, WHITE)
            set_cell_border(cell, TABLE_RULE, "5")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=(first_col_bold and col_idx == 0))
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_paragraph(doc, text, *, bold_lead=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, size=11):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, size=size, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=size, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_number(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True)
    return p


def add_placeholder(doc, title, instruction, height_lines=5):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [16.8])
    cell = table.cell(0, 0)
    shade_cell(cell, WHITE)
    set_cell_border(cell, TABLE_RULE, "8", "dashed")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(max(12, height_lines * 8))
    r2 = p2.add_run(instruction)
    set_run_font(r2, size=9, italic=True)
    return table


def add_callout(doc, label, text, color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [16.8])
    cell = table.cell(0, 0)
    shade_cell(cell, WHITE)
    set_cell_border(cell, TABLE_RULE, "6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_run_font(r, size=9, bold=True)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_break(doc):
    doc.add_page_break()


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(10.5)
    specs = {
        "Heading 1": (15, BLACK, 10, 5),
        "Heading 2": (12.5, BLACK, 8, 4),
        "Heading 3": (11.5, BLACK, 6, 3),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)
        section.different_first_page_header_footer = True
        # Tidak ada running header sesuai arahan desain sederhana.
        header = section.header
        p = header.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = Pt(0)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(fp)


def make_logo_placeholder(cell, text):
    shade_cell(cell, WHITE)
    set_cell_border(cell, TABLE_RULE, "6", "dashed")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, name="Arial", size=7.5, bold=True)


def add_cover(doc):
    logos = doc.add_table(rows=1, cols=3)
    set_table_widths(logos, [5.6, 5.6, 5.6])
    make_logo_placeholder(logos.cell(0, 0), "TEMPAT LOGO\nKEMENDIKTISAINTEK\n(pojok kiri atas)")
    shade_cell(logos.cell(0, 1), WHITE)
    set_cell_border(logos.cell(0, 1), WHITE, "0", "nil")
    make_logo_placeholder(logos.cell(0, 2), "TEMPAT LOGO GEMASTIK XIX 2026\nDAN UNIVERSITAS GUNADARMA\n(berdampingan di pojok kanan atas)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROPOSAL BABAK PENYISIHAN")
    set_run_font(r, name="Arial", size=12, bold=True)
    p.paragraph_format.space_after = Pt(4)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("PAGELARAN MAHASISWA NASIONAL BIDANG TIK\nGEMASTIK XIX TAHUN 2026")
    set_run_font(r, name="Arial", size=10.5, bold=True)
    p2.paragraph_format.space_after = Pt(15)
    logo = MEDIA / "image3.jpg"
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(p3.add_run(), logo, Cm(7.0), "Logo REKSA, Radar Evaluasi K3 dan Sensor Ancaman Kerja")
    p3.paragraph_format.space_after = Pt(8)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_after = Pt(8)
    r = p4.add_run("RADAR EVALUASI K3 DAN SENSOR ANCAMAN KERJA")
    set_run_font(r, name="Arial", size=18, bold=True)
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_after = Pt(16)
    r = p5.add_run(
        "Smart Helmet Offline-First Berbasis AI-IoT untuk Peringatan Zona Bahaya Dinamis "
        "dan Intelijen Near-Miss pada Lingkungan Industri"
    )
    set_run_font(r, name="Arial", size=12.5, italic=True)
    add_table(
        doc,
        ["IDENTITAS TIM", "KETERANGAN"],
        [
            ("Nama Tim", "Restu Bundo"),
            ("Ketua", "Wiefran Varenzo"),
            ("Anggota", "Felix Rafael; Muhammad Zaki Alfadilah"),
            ("Dosen Pendamping", "Dr. Guntur Eka Saputra, S.T., M.M.S.I."),
            ("Perguruan Tinggi", "Universitas Gunadarma"),
            ("Kategori", "Piranti Cerdas, Sistem Benam & IoT"),
        ],
        [5.0, 11.8],
        font_size=9.6,
        first_col_bold=True,
    )
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(8)
    r = p6.add_run("UNIVERSITAS GUNADARMA\n2026")
    set_run_font(r, name="Arial", size=10.5, bold=True)
    add_page_break(doc)


def add_toc(doc):
    add_heading(doc, "DAFTAR ISI", 1)
    entries = [
        ("1. Abstrak", "3"),
        ("2. Latar Belakang", "4"),
        ("3. Urgensi dan Manfaat", "5"),
        ("3.1 Urgensi Karya dan Tujuan Pengembangan", "5"),
        ("3.2 Manfaat bagi Pemangku Kepentingan", "5"),
        ("3.3 Keunggulan dan Kebaruan Karya", "6"),
        ("3.4 Perbandingan dengan Karya Sejenis", "6"),
        ("4. Metode Dasar Pengembangan Karya", "7"),
        ("4.1 Penerapan Tiga Elemen Teknologi", "7"),
        ("5. Desain Purwarupa / Model", "8"),
        ("5.1 Arsitektur Sistem Offline-First", "8"),
        ("5.2 Desain Sistem Benam", "9"),
        ("5.5 Komunikasi IoT dan Cloud", "11"),
        ("5.6 Infrastruktur Cloud", "12"),
        ("5.7 AI Priority Engine dan Early-Warning Advisory", "13"),
        ("6. Analisis Fungsional, Cara Kerja, dan Kinerja", "14"),
        ("7. Rencana Implementasi dan Perkembangan Pengerjaan", "17"),
        ("8. Foto dan Penjelasan Hasil Implementasi Purwarupa", "20"),
        ("9. Tautan Video Proses Pengembangan Model Karya Inovasi", "23"),
        ("10. Daftar Pustaka", "24"),
    ]
    table = doc.add_table(rows=0, cols=2)
    set_table_widths(table, [14.8, 2.0])
    for text, page in entries:
        cells = table.add_row().cells
        for c in cells:
            set_cell_border(c, WHITE, "0", "nil")
            set_cell_margins(c, top=30, bottom=30)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=(text[0].isdigit() and "." in text[:3] and text.count(".") == 1))
        p2 = cells[1].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(page)
        set_run_font(r2, size=10.5, color=GRAY)
    add_page_break(doc)


def add_abstract(doc):
    add_heading(doc, "1. ABSTRAK", 1)
    add_paragraph(
        doc,
        "Keselamatan pada warehouse dan area intralogistik tidak hanya ditentukan oleh keberadaan alat pelindung diri, tetapi juga oleh kemampuan pekerja mengenali hazard bergerak pada saat yang tepat. Forklift, troli bermotor, mesin aktif, dan lorong sempit membentuk zona risiko yang berubah mengikuti aktivitas operasional. Di sisi lain, kejadian nyaris celaka sering tidak meninggalkan bukti yang cukup untuk ditinjau kembali karena pencatatannya masih bergantung pada laporan manual. REKSA, singkatan dari Radar Evaluasi K3 dan Sensor Ancaman Kerja, dikembangkan sebagai smart helmet berbasis AI-IoT yang memberi peringatan lokal sekaligus membentuk data keselamatan yang dapat diaudit.",
    )
    add_paragraph(
        doc,
        "Sistem terdiri atas Active Hazard Node berbasis ESP32 yang memancarkan identitas hazard melalui Bluetooth Low Energy dan Smart Helmet Node berbasis ESP32 yang memindai sinyal tersebut. Helmet node memproses RSSI dengan penyaringan, hysteresis, dan time persistence untuk membentuk state aman, waspada, dan kritis. Keputusan alarm dijalankan langsung pada perangkat melalui buzzer dan vibration motor. Dengan rancangan offline-first ini, fungsi peringatan tetap aktif saat Wi-Fi, broker, backend, layanan AI, atau internet tidak tersedia. MPU6050 merekam konteks gerakan dan kandidat benturan atau jatuh, sedangkan DHT22 serta MQ135 memberikan konteks suhu, kelembapan, dan perubahan relatif kualitas udara.",
    )
    add_paragraph(
        doc,
        "Telemetry dikirim menggunakan MQTT QoS 1 melalui koneksi TLS ke HiveMQ Cloud. Firmware menyimpan data yang belum diakui pada buffer, melakukan retransmisi, dan baru menghapusnya setelah menerima ACK. Gateway bridge menyimpan pesan pada outbox SQLite sebelum meneruskannya ke Google Cloud Pub/Sub, lalu menjalankan deduplikasi persisten berdasarkan message_id. Backend FastAPI dan layanan AI berjalan pada Cloud Run, raw envelope diarsipkan pada Firestore, dan dashboard supervisor dipublikasikan melalui Firebase Hosting. Uji integrasi perangkat nyata telah membuktikan alur ESP32 hingga dashboard cloud, termasuk respons HTTP 204 pada endpoint Pub/Sub dan pembaruan data pekerja W01 dari sumber HARDWARE_MQTT.",
    )
    add_paragraph(
        doc,
        "Pada tahap sekarang, komponen kecerdasan dijalankan sebagai RULE_BASED_BASELINE dengan peran ADVISORY_ONLY. REKSA tidak mengizinkan model AI mengendalikan alarm lokal. Dataset perangkat nyata akan disusun dari skenario terkontrol dan diverifikasi per sesi untuk membandingkan baseline aturan, Logistic Regression, dan Random Forest. Model hanya akan digunakan apabila mengungguli baseline pada recall, F1-score, false alarm, calibration, dan held-out participant. Pendekatan ini menjadikan REKSA bukan sekadar helm bersensor, melainkan sistem keselamatan berlapis yang tetap responsif di lapangan dan sekaligus menghasilkan dasar evaluasi near-miss yang dapat dipertanggungjawabkan.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Kata kunci: ")
    set_run_font(r, bold=True, size=10.5)
    r2 = p.add_run("AI-IoT, smart helmet, BLE proximity, offline-first safety, MQTT QoS 1, near-miss intelligence, Cloud Run.")
    set_run_font(r2, italic=True, size=10.5)
    add_page_break(doc)


def add_background(doc):
    add_heading(doc, "2. LATAR BELAKANG", 1)
    add_paragraph(
        doc,
        "Kecelakaan kerja masih menjadi persoalan yang memerlukan perhatian serius. BPJS Ketenagakerjaan mencatat sekitar 141 ribu kasus kecelakaan kerja sepanjang 2024, sementara Program Nasional Keselamatan dan Kesehatan Kerja 2024-2029 menempatkan penguatan budaya pencegahan sebagai agenda penting di Indonesia (BPJS Ketenagakerjaan, 2025; ILO, 2024). Angka tersebut tidak seluruhnya berasal dari warehouse, tetapi cukup menunjukkan bahwa intervensi K3 tidak dapat hanya berhenti pada penanganan setelah kecelakaan. Area gudang dan intralogistik mempunyai tantangan khusus karena pekerja berjalan kaki berbagi ruang dengan kendaraan material handling, rak logam, mesin, dan jalur operasional yang berubah sepanjang hari.",
    )
    add_paragraph(
        doc,
        "Marka lantai, cermin tikungan, SOP, dan rompi reflektif tetap merupakan pengendalian utama yang tidak boleh digantikan. Masalah muncul ketika hazard bergerak tidak mempunyai zona bahaya yang tetap. Pekerja dapat membelakangi kendaraan, pandangan dapat tertutup rak, dan kebisingan dapat mengurangi efektivitas alarm suara. Karena itu, dibutuhkan lapisan peringatan personal yang berada dekat dengan pengguna dan mampu merespons tanpa menunggu jaringan. Kajian smart helmet menunjukkan bahwa wearable telah digunakan untuk pemantauan aktivitas, kondisi lingkungan, dan peringatan bahaya. Kim, Baek, dan Choi (2021) juga menunjukkan bahwa beacon Bluetooth pada kendaraan dan area berbahaya dapat diterima oleh smart helmet untuk membentuk peringatan proksimitas.",
    )
    add_paragraph(
        doc,
        "Bluetooth Low Energy dipilih karena telah terintegrasi pada ESP32, konsumsi dayanya relatif rendah, dan advertising dapat diterima tanpa proses pairing. Walaupun demikian, RSSI BLE tidak dapat diperlakukan sebagai jarak absolut. Orientasi antena, tubuh manusia, rak logam, multipath, dan kondisi ruangan dapat mengubah pembacaan pada posisi yang sama. REKSA tidak mengubah RSSI menjadi klaim meter. Nilai tersebut dipakai sebagai indikator relatif yang terlebih dahulu disaring, kemudian diproses oleh hysteresis dan time persistence. Dengan cara ini, state tidak mudah berosilasi akibat satu sampel yang menyimpang.",
    )
    add_paragraph(
        doc,
        "Persoalan kedua adalah hilangnya informasi sebelum kecelakaan terjadi. Near-miss penting karena menunjukkan keberadaan hazard yang berpotensi menghasilkan konsekuensi lebih besar. Studi prospektif Inagaki et al. (2024) menemukan hubungan antara kecukupan respons perusahaan terhadap laporan near-miss dan kejadian kecelakaan pada periode berikutnya. Namun, pelaporan cedera maupun kejadian terkait kerja dapat terhambat oleh kekhawatiran terhadap konsekuensi pekerjaan, beban administrasi, keterbatasan pengetahuan, dan anggapan bahwa kejadian tidak cukup serius untuk dilaporkan (Kyung et al., 2023). Artinya, sistem perlu membantu pencatatan tanpa langsung mengambil alih penilaian supervisor.",
    )
    add_paragraph(
        doc,
        "REKSA menjawab dua kebutuhan tersebut melalui pemisahan jalur keselamatan dan jalur analitik. Jalur pertama berada pada helmet node dan bertanggung jawab atas alarm lokal. Jalur kedua mengirim telemetry serta event ke cloud untuk dashboard, audit, dan pengembangan AI. Event otomatis tidak langsung diberi label near-miss. Sistem membedakan unsafe proximity event, candidate near-miss, dan verified near-miss. Verifikasi akhir tetap dilakukan supervisor K3 sehingga data yang masuk ke dataset mempunyai asal label yang jelas.",
    )
    add_callout(
        doc,
        "Rumusan masalah",
        "Bagaimana merancang piranti wearable yang mampu memberi peringatan terhadap hazard dinamis secara lokal ketika konektivitas gagal, sekaligus mencatat bukti kejadian secara andal dan memanfaatkan AI secara terukur tanpa menjadikannya pengendali keselamatan?",
        color=RED,
    )
    add_page_break(doc)


def add_urgency(doc):
    add_heading(doc, "3. URGENSI DAN MANFAAT", 1)
    add_heading(doc, "3.1 Urgensi Karya dan Tujuan Pengembangan", 2)
    add_paragraph(
        doc,
        "Urgensi REKSA terletak pada celah antara peringatan lapangan dan evaluasi setelah kejadian. Perangkat komersial tertentu telah menawarkan proximity warning, sedangkan dashboard IoT dapat menyajikan telemetry. Namun, purwarupa pendidikan sering menempatkan seluruh keputusan pada server atau menganggap setiap nilai sensor yang melewati ambang sebagai near-miss. REKSA mengambil posisi yang lebih hati-hati. Alarm wajib tetap bekerja di perangkat, pengiriman data harus tahan gangguan, dan hasil AI harus dapat ditelusuri sebagai rekomendasi, bukan keputusan final.",
    )
    add_table(
        doc,
        ["Kesenjangan", "Tujuan Pengembangan", "Bukti Implementasi Saat Ini"],
        [
            ("Hazard bergerak tidak mempunyai zona tetap.", "Membangun Active Hazard Zone berbasis BLE yang mengikuti identitas hazard.", "Hazard node F01 terdeteksi oleh helmet W01 dan memicu transisi SAFE hingga CRITICAL."),
            ("Alarm berbasis jaringan rentan terlambat atau tidak tersedia.", "Menjalankan klasifikasi zona dan aktuator langsung pada ESP32.", "Buzzer dan vibration motor tetap bekerja ketika MQTT belum terhubung."),
            ("Telemetry dapat hilang saat koneksi putus.", "Menggunakan QoS 1, buffer, retransmisi, ACK, dan deduplikasi persisten.", "Message_id yang diulang hanya di-ACK kembali dan tidak diteruskan dua kali ke Pub/Sub."),
            ("Near-miss sering tidak tercatat atau langsung diklaim otomatis.", "Membangun taksonomi event dan alur verifikasi supervisor.", "Backend membedakan telemetry, candidate event, dan status verifikasi."),
            ("AI sering ditempelkan tanpa pembanding dan batas peran.", "Menguji model terhadap baseline dan menempatkannya sebagai advisory-only.", "Layanan AI menampilkan mode RULE_BASED_BASELINE serta provenance hasil."),
        ],
        [4.3, 6.1, 6.4],
        font_size=8.7,
    )
    add_heading(doc, "3.2 Manfaat bagi Pemangku Kepentingan", 2)
    add_table(
        doc,
        ["Pemangku Kepentingan", "Manfaat Utama"],
        [
            ("Pekerja", "Menerima peringatan audio dan haptik ketika memasuki zona risiko, termasuk saat internet terputus."),
            ("Supervisor K3", "Melihat status pekerja, konteks sensor, riwayat paparan, dan antrean candidate near-miss untuk ditinjau."),
            ("Manajemen Operasional", "Memperoleh data frekuensi, durasi, dan pola paparan sebagai bahan evaluasi tata letak serta prosedur kerja."),
            ("Tim Teknis", "Mempunyai jejak message_id, log cloud, dan dataset terstruktur untuk pemeliharaan serta pengembangan model."),
            ("Institusi Pendidikan", "Memperoleh contoh integrasi AI, embedded system, dan IoT yang dapat diuji, diukur, dan direplikasi."),
        ],
        [4.8, 12.0],
        font_size=9.2,
        first_col_bold=True,
    )
    add_paragraph(
        doc,
        "Dari sisi keberlanjutan, REKSA berkaitan langsung dengan SDG 8, khususnya perlindungan lingkungan kerja yang aman, serta SDG 9 melalui pemanfaatan infrastruktur digital dan inovasi industri. Kontribusi tersebut tidak dinyatakan sebagai dampak final, melainkan sebagai arah manfaat yang akan diuji melalui pengurangan missed warning, kestabilan alarm, kualitas data event, dan kemudahan tindak lanjut supervisor.",
    )
    add_paragraph(
        doc,
        "Potensi kegunaan bagi masyarakat tidak berhenti pada satu perusahaan. Arsitektur REKSA dapat direplikasi pada gudang skala kecil, laboratorium pendidikan, bengkel, dan area produksi yang mempunyai interaksi antara pekerja dengan hazard bergerak. Biaya node yang relatif terjangkau, mekanisme alarm yang tetap bekerja tanpa internet, serta penggunaan platform umum membuat penerapannya dapat dimulai secara bertahap. Data yang terkumpul juga dapat membantu organisasi menyusun intervensi berbasis bukti, misalnya memperbaiki jalur material handling, titik buta, atau prosedur kerja pada area yang berulang kali menghasilkan candidate near-miss.",
    )
    add_heading(doc, "3.3 Keunggulan dan Kebaruan Karya", 2)
    innovations = [
        ("Offline-first safety", "Alarm ditentukan oleh state machine ESP32. Cloud dan AI tidak berada pada jalur kritis."),
        ("Active Hazard Zone", "Hazard tetap dan bergerak memakai identitas BLE yang sama-sama dapat diikuti helmet."),
        ("Delivery yang dapat diaudit", "QoS 1 dilengkapi message_id, buffer, ACK, retransmisi, durable outbox, dan deduplikasi."),
        ("Taksonomi near-miss bertingkat", "Sistem membedakan pembacaan sensor, kandidat kejadian, dan hasil verifikasi supervisor."),
        ("AI dengan evidence gate", "Model hanya dipromosikan bila mengungguli baseline pada evaluasi berbasis sesi dan peserta."),
        ("Cloud sebagai penguat, bukan ketergantungan", "HiveMQ, Pub/Sub, Cloud Run, Firestore, dan Firebase memperkuat audit serta akses tanpa mematikan alarm saat gagal."),
    ]
    for title, body in innovations:
        add_paragraph(doc, f"{title}. {body}", bold_lead=f"{title}.", after=4, size=10.5)
    add_heading(doc, "3.4 Perbandingan dengan Karya Sejenis", 2)
    add_table(
        doc,
        ["Dimensi", "Smart Helmet Proximity", "Dashboard IoT", "REKSA"],
        [
            ("Alarm lokal tanpa internet", "Bervariasi", "Tidak", "Ya, fungsi inti"),
            ("Hazard bergerak sebagai zona aktif", "Sebagian", "Bergantung input", "Ya"),
            ("Buffer, ACK, dan deduplikasi", "Jarang dijelaskan", "Bervariasi", "Ya, end-to-end"),
            ("Taksonomi near-miss dan verifikasi", "Umumnya tidak", "Sebagian", "Ya"),
            ("AI dibandingkan dengan baseline", "Tidak selalu", "Tidak selalu", "Wajib sebelum deploy"),
            ("Provenance model di dashboard", "Tidak", "Tidak selalu", "Ya"),
            ("Cloud tidak mengendalikan alarm", "Bervariasi", "Tidak relevan", "Ya"),
        ],
        [3.8, 4.1, 4.0, 4.9],
        font_size=8.5,
    )
    add_paragraph(
        doc,
        "Perbandingan ini menjelaskan posisi arsitektural, bukan klaim bahwa seluruh produk di pasar mempunyai karakteristik yang sama. Rujukan smart helmet BLE dari Kim, Baek, dan Choi (2021) memperlihatkan kelayakan konsep proximity warning. Kebaruan REKSA berada pada integrasi alarm offline-first, ketahanan delivery, taksonomi event, dan evidence gate untuk AI dalam satu alur purwarupa.",
        size=10.2,
    )
    add_page_break(doc)


def add_method(doc):
    add_heading(doc, "4. METODE DASAR PENGEMBANGAN KARYA", 1)
    add_paragraph(
        doc,
        "Pengembangan menggunakan pendekatan iterative engineering. Setiap iterasi dimulai dari risiko keselamatan yang ingin dikurangi, diterjemahkan menjadi fungsi yang dapat diuji, lalu dibuktikan melalui log perangkat dan layanan. Pendekatan ini dipilih agar penambahan fitur tidak mengaburkan fungsi inti dan agar klaim proposal selalu sejalan dengan kondisi purwarupa.",
    )
    stages = [
        ("Studi literatur dan konteks K3", "Mengkaji near-miss, underreporting, smart helmet, BLE proximity, MQTT, dan prinsip manajemen K3. Hasil studi dipakai untuk menentukan batas klaim dan kebutuhan verifikasi manusia."),
        ("Analisis kebutuhan dan failure mode", "Memetakan pengguna, hazard, kondisi jaringan, kegagalan daya, fluktuasi RSSI, dan risiko data duplikat. Alarm lokal ditetapkan sebagai fungsi yang tidak boleh bergantung pada komponen eksternal."),
        ("Perancangan arsitektur berlapis", "Memisahkan Local Safety Layer, Reliable IoT Layer, Cloud Analytics Layer, dan AI Advisory Layer. Setiap batas layanan mempunyai kontrak data dan mekanisme kegagalan yang jelas."),
        ("Pengembangan sistem benam", "Merakit helmet node serta hazard node, mengintegrasikan BLE, MPU6050, DHT22, MQ135, buzzer, vibration motor, dan driver aktuator."),
        ("Pengembangan firmware offline-first", "Menerapkan pemindaian BLE, smoothing RSSI, hysteresis, time persistence, state machine alarm, buffer 12 record, retransmisi tiga detik, NTP untuk TLS, dan ACK."),
        ("Pengembangan IoT dan cloud", "Menggunakan HiveMQ Cloud TLS 8883, MQTT QoS 1, gateway bridge dengan outbox SQLite, Pub/Sub push, Cloud Run, Firestore, dan Firebase Hosting."),
        ("Pengembangan kecerdasan", "Membangun rule-based baseline, mendefinisikan fitur temporal, dan menyiapkan perbandingan Logistic Regression serta Random Forest dengan grouped holdout."),
        ("Integrasi dan verifikasi", "Menguji koneksi, health endpoint, forwarding message_id, deduplikasi, status dashboard, serta perilaku alarm ketika jaringan tidak tersedia."),
    ]
    add_table(doc, ["Tahap", "Pelaksanaan pada REKSA"], stages, [4.3, 12.5], font_size=9.1, first_col_bold=True)
    add_heading(doc, "4.1 Penerapan Tiga Elemen Teknologi", 2)
    add_table(
        doc,
        ["Elemen", "Penerapan", "Bukti yang Dapat Ditunjukkan kepada Juri"],
        [
            ("Kecerdasan", "Priority engine dan early-warning advisory berbasis fitur temporal. Saat ini baseline aturan; kandidat ML menunggu dataset nyata.", "Endpoint /health menampilkan mode model, safety_role ADVISORY_ONLY, provenance, dan versi artifact."),
            ("Sistem Benam", "ESP32 memproses BLE dan sensor serta mengendalikan buzzer dan motor getar secara lokal.", "Serial monitor menunjukkan transisi zona dan ALARM_APPLIED ketika MQTT belum terhubung."),
            ("Internet of Things", "MQTT TLS dan QoS 1 ke HiveMQ, bridge durable ke Pub/Sub, Cloud Run, Firestore, serta dashboard Firebase.", "Log message_id dari Buffered, Forwarded, POST 204, hingga pembaruan W01 pada dashboard."),
        ],
        [3.2, 7.0, 6.6],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Prinsip pengembangan",
        "Fitur baru tidak dinilai dari jumlah teknologinya, tetapi dari relevansi terhadap masalah, keterukuran, ketahanan saat gagal, dan kemampuan tim menjelaskan bukti implementasinya.",
        color=GREEN,
    )
    add_page_break(doc)


def add_design(doc):
    add_heading(doc, "5. DESAIN PURWARUPA / MODEL", 1)
    add_heading(doc, "5.1 Arsitektur Sistem Offline-First", 2)
    add_paragraph(
        doc,
        "REKSA menggunakan empat lapisan yang mempunyai tanggung jawab berbeda. Local Safety Layer berada pada helmet dan menjadi satu-satunya lapisan yang boleh mengaktifkan alarm keselamatan tanpa ketergantungan eksternal. Reliable IoT Layer memastikan telemetry tidak langsung hilang ketika jaringan berubah. Cloud Analytics Layer mengolah dan menyajikan data. AI Advisory Layer memberi prioritas pemeriksaan serta estimasi risiko secara nonkritis.",
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 1: ARSITEKTUR END-TO-END REKSA",
        "Gambarkan alur berwarna: Active Hazard Node -> Smart Helmet (alarm lokal) -> HiveMQ Cloud TLS/QoS 1 -> Gateway Outbox SQLite -> Pub/Sub -> Cloud Run Backend + Cloud Run AI -> Firestore -> Firebase Dashboard. Gunakan garis tebal merah hanya untuk jalur alarm lokal dan garis biru putus-putus untuk telemetry. Tulis bahwa kegagalan cloud tidak mematikan alarm.",
        7,
    )
    add_caption(doc, "Gambar 1. Arsitektur berlapis REKSA dan pemisahan jalur keselamatan dari jalur cloud.")
    add_table(
        doc,
        ["Lapisan", "Fungsi", "Respons ketika Komponen Gagal"],
        [
            ("Local Safety", "Deteksi hazard, klasifikasi zona, buzzer, vibration motor.", "Tetap berjalan selama helmet memperoleh daya."),
            ("Reliable IoT", "Buffer, ACK, retransmisi, deduplikasi, dan forwarding.", "Record ditahan dan dikirim ulang ketika koneksi pulih."),
            ("Cloud Analytics", "Ingest, arsip, API, WebSocket, dan dashboard.", "Monitoring tertunda; alarm lokal tidak berubah."),
            ("AI Advisory", "Skor prioritas dan early-warning advisory.", "Status UNAVAILABLE; state machine lokal tetap berlaku."),
        ],
        [3.4, 7.0, 6.4],
        font_size=9.1,
    )
    add_page_break(doc)
    add_heading(doc, "5.2 Desain Sistem Benam", 2)
    add_heading(doc, "Smart Helmet Node", 3)
    add_paragraph(
        doc,
        "Smart Helmet Node memakai ESP32 sebagai pengendali utama. BLE radio memindai advertising hazard, Wi-Fi mengirim telemetry, dan mikrokontroler menjalankan state machine alarm. MPU6050 menyediakan percepatan serta kecepatan sudut. DHT22 mencatat suhu dan kelembapan, sedangkan MQ135 dipakai sebagai indikator relatif kualitas udara setelah warm-up dan baseline. Buzzer serta vibration motor dikendalikan melalui driver agar beban tidak ditarik langsung dari GPIO.",
    )
    add_table(
        doc,
        ["Komponen", "Antarmuka / Pin", "Peran", "Batasan"],
        [
            ("ESP32 Dev Module", "Wi-Fi + BLE", "Pemrosesan lokal dan komunikasi", "Purwarupa, belum tersertifikasi"),
            ("MPU6050", "I2C SDA 32, SCL 33", "Konteks gerakan dan candidate impact/fall", "Bukan diagnosis cedera"),
            ("DHT22", "GPIO 18", "Suhu dan kelembapan", "Bukan WBGT"),
            ("MQ135", "ADC GPIO 34", "Perubahan relatif kualitas udara", "Bukan AQI atau gas spesifik"),
            ("Active buzzer", "GPIO 19 melalui driver", "Alarm audio", "Efektivitas dipengaruhi kebisingan"),
            ("Vibration motor", "GPIO 23 melalui driver", "Alarm haptik", "Dapat memengaruhi pembacaan IMU"),
        ],
        [3.5, 3.6, 5.2, 4.5],
        font_size=8.4,
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 2: SKEMA RANGKAIAN DAN DISTRIBUSI DAYA",
        "Masukkan diagram wiring aktual. Tampilkan ESP32, MPU6050, DHT22, MQ135, transistor/MOSFET driver, flyback protection untuk motor, sumber 5 V, regulator 3,3 V, ground bersama, dan batas tegangan ADC. Jangan memakai diagram generik yang tidak sama dengan rangkaian fisik.",
        5,
    )
    add_caption(doc, "Gambar 2. Rangkaian Smart Helmet Node dan pemisahan jalur sensor dari beban aktuator.")
    add_heading(doc, "Active Hazard Node", 3)
    add_paragraph(
        doc,
        "Hazard node menggunakan ESP32 sebagai BLE advertiser dengan nama terstruktur, misalnya REKSA_HAZARD_F01. Node dapat ditempatkan pada miniatur forklift, troli, mesin, atau area terbatas. Pada purwarupa saat ini, helmet memvalidasi prefix identitas dan memelihara daftar hazard aktif berdasarkan waktu terakhir terlihat. Dengan demikian, hazard yang keluar dari jangkauan tidak terus diperlakukan aktif.",
    )
    add_page_break(doc)
    add_heading(doc, "5.3 State Machine Zona dan Alarm", 2)
    add_paragraph(
        doc,
        "Klasifikasi zona memakai RSSI rata-rata, arah transisi, dan lama kondisi bertahan. Hysteresis membuat ambang masuk dan keluar berbeda, sedangkan time persistence mencegah satu pembacaan ekstrem langsung mengubah state. Hazard dengan tingkat paling berbahaya dipilih sebagai sumber peringatan prioritas. Semua proses tersebut berlangsung sebelum telemetry dikirim.",
    )
    add_table(
        doc,
        ["State", "Interpretasi", "Respons Lokal", "Catatan Event"],
        [
            ("SAFE", "Belum memenuhi kondisi risiko", "Buzzer dan motor mati", "Telemetry periodik"),
            ("MODERATE", "Pendekatan mulai perlu diperhatikan", "Getaran intermiten", "Durasi mulai diamati"),
            ("HIGH", "Paparan meningkat", "Getaran cepat dan buzzer berkala", "Unsafe proximity event"),
            ("CRITICAL", "Zona terdekat atau kombinasi risiko tinggi", "Buzzer dan motor menyala terus", "Candidate near-miss jika syarat konteks terpenuhi"),
            ("RECOVERY", "Kondisi aman bertahan setelah keluar zona", "Alarm dihentikan", "Event ditutup dan durasi dihitung"),
        ],
        [2.6, 5.0, 4.3, 4.9],
        font_size=8.8,
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 3: STATE MACHINE ZONA",
        "Buat diagram state SAFE -> MODERATE -> HIGH -> CRITICAL dan jalur recovery. Cantumkan ambang hasil kalibrasi aktual, nilai hysteresis, time persistence, serta kondisi hazard timeout. Jangan mengisi angka sebelum hasil kalibrasi final tersedia.",
        5,
    )
    add_caption(doc, "Gambar 3. State machine zona dengan filtering, hysteresis, dan time persistence.")
    add_heading(doc, "5.4 Sumber Daya dan Keselamatan Kelistrikan", 2)
    add_paragraph(
        doc,
        "Purwarupa saat ini paling stabil pada sumber 5 V USB atau power bank. Pengujian baterai harus memperhitungkan arus puncak Wi-Fi, heater MQ135, buzzer, dan motor getar. Sebelum pemasangan akhir, tim akan mengukur arus idle, arus rata-rata, arus puncak, durasi operasi, dan penurunan tegangan. Kabel serta modul ditempatkan pada casing yang tidak mengganggu struktur helm dan tidak mempunyai bagian konduktif terbuka.",
    )
    add_callout(doc, "Batas keselamatan", "REKSA adalah purwarupa lapisan peringatan tambahan. Sistem tidak menggantikan helm bersertifikat, SOP, marka jalur, spotter, inspeksi kendaraan, atau pengendalian K3 yang telah berlaku.", color=RED)
    add_page_break(doc)
    add_heading(doc, "5.5 Komunikasi IoT dan Cloud", 2)
    add_paragraph(
        doc,
        "Helmet terhubung ke HiveMQ Cloud menggunakan MQTT melalui TLS pada port 8883. Kredensial dipisahkan untuk helmet, backend lokal, dan gateway. Payload memakai timestamp UTC serta message_id unik. QoS 1 memastikan broker memberikan pengakuan pada level protokol, sedangkan ACK aplikasi memastikan record benar-benar diterima backend atau bridge. Pemisahan ini penting karena PUBACK dari broker belum berarti data telah masuk ke pipeline cloud.",
    )
    add_table(
        doc,
        ["Kontrak", "Implementasi"],
        [
            ("Topic telemetry", "REKSA/helmet/W01/sensor"),
            ("Topic warning", "REKSA/helmet/W01/warning"),
            ("Topic ACK", "REKSA/helmet/W01/telemetry/ack"),
            ("Transport", "MQTT QoS 1 melalui TLS 1.2+, port 8883"),
            ("Identitas record", "message_id unik per telemetry"),
            ("Buffer helmet", "Maksimum 12 payload di RAM; record tertua diganti jika penuh"),
            ("Retry", "Setiap tiga detik hingga ACK diterima"),
            ("Gateway durability", "SQLite outbox sebelum publish ke Pub/Sub"),
            ("Deduplication", "Riwayat delivered message_id persisten; duplikat hanya di-ACK ulang"),
        ],
        [5.0, 11.8],
        font_size=9.1,
        first_col_bold=True,
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 4: SEQUENCE DELIVERY DAN ACK",
        "Gambarkan urutan Helmet publish -> HiveMQ PUBACK -> Bridge simpan outbox -> Pub/Sub message ID -> Bridge tandai delivered -> ACK aplikasi -> Helmet hapus buffer. Tambahkan jalur putus koneksi dan retransmisi. Gunakan satu contoh message_id nyata yang sudah disamarkan bila diperlukan.",
        6,
    )
    add_caption(doc, "Gambar 4. Mekanisme at-least-once delivery, ACK aplikasi, dan deduplikasi REKSA.")
    add_page_break(doc)
    add_heading(doc, "5.6 Infrastruktur Cloud", 2)
    add_paragraph(
        doc,
        "Gateway meneruskan envelope MQTT ke topic Pub/Sub reksa-telemetry. Push subscription yang diautentikasi memanggil endpoint ingest pada backend Cloud Run. Backend memvalidasi envelope, memperbarui state dashboard, memanggil layanan AI secara advisory, dan menyimpan raw envelope ke Firestore untuk audit. Frontend React dibangun dengan URL backend cloud dan dipublikasikan melalui Firebase Hosting. Min-instances diatur nol agar biaya purwarupa tetap rendah ketika tidak digunakan.",
    )
    add_table(
        doc,
        ["Layanan", "Peran", "Status Implementasi"],
        [
            ("HiveMQ Cloud Serverless", "Broker MQTT TLS", "Berjalan"),
            ("Gateway bridge", "MQTT ke Pub/Sub + durable outbox", "Berjalan pada laptop/gateway demo"),
            ("Pub/Sub", "Antrian asinkron", "Topic dan push subscription aktif"),
            ("Cloud Run Backend", "API, ingest, state, dan arsip", "Healthy"),
            ("Cloud Run AI", "Priority dan forecaster advisory", "Healthy, rule-based baseline"),
            ("Firestore", "Arsip raw envelope", "Aktif"),
            ("Firebase Hosting", "Dashboard supervisor", "Aktif"),
        ],
        [4.2, 7.0, 5.6],
        font_size=8.9,
    )
    add_paragraph(doc, "Dashboard publik purwarupa: https://reksa-505809.web.app", bold_lead="Dashboard publik purwarupa:", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.3)
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 5: BUKTI DEPLOYMENT CLOUD",
        "Gabungkan empat tangkapan layar yang diberi label: HiveMQ connected, bridge Forwarded ke Pub/Sub, Cloud Run POST /api/v1/ingest/pubsub 204, dan dashboard Firebase yang menampilkan W01. Samarkan credential, token, project secret, dan service-account key.",
        6,
    )
    add_caption(doc, "Gambar 5. Bukti integrasi perangkat hingga dashboard cloud.")
    add_page_break(doc)
    add_heading(doc, "5.7 AI Priority Engine dan Early-Warning Advisory", 2)
    add_paragraph(
        doc,
        "Kecerdasan REKSA diarahkan pada dua tugas yang saling melengkapi. Priority Engine mengurutkan candidate event agar supervisor memeriksa kejadian paling mendesak terlebih dahulu. Early-Warning Advisory menggunakan window temporal RSSI dan gerakan untuk memperkirakan pola pendekatan sebelum state lokal menjadi kritis. Keduanya bersifat advisory-only. Jika layanan AI lambat atau gagal, backend menampilkan status UNAVAILABLE dan tidak menghasilkan probabilitas pengganti.",
    )
    add_table(
        doc,
        ["Engine", "Fitur Kandidat", "Target", "Keputusan Saat Ini"],
        [
            ("Priority Engine", "Durasi paparan, RSSI terfilter, state hazard, impact/fall candidate, suhu, kualitas udara, riwayat event", "Urgensi pemeriksaan supervisor", "Baseline aturan"),
            ("Early-Warning Forecaster", "Slope RSSI, variance, dwell time, acceleration, gyroscope, trajectory window", "Advisory pendekatan 3-5 detik", "Baseline aturan"),
        ],
        [3.5, 6.5, 3.9, 2.9],
        font_size=8.5,
    )
    add_paragraph(
        doc,
        "Tahap pembelajaran mesin dimulai setelah dataset perangkat nyata cukup. Data akan direkam pada 5-10 Hz dengan participant_id, session_id, scenario, dan trajectory_id. Pembagian train-test dilakukan berdasarkan kelompok participant dan session agar sampel dari satu rangkaian gerak tidak bocor ke kedua sisi. Baseline aturan, Logistic Regression, dan Random Forest dibandingkan pada recall, precision, F1-score, false alarm per jam, Brier score, confusion matrix, dan calibration. Model tidak dipromosikan hanya karena akurasi totalnya tinggi.",
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 6: PIPELINE DATA DAN PERBANDINGAN MODEL",
        "Tampilkan alur raw telemetry -> temporal window -> feature engineering -> grouped holdout -> rule-based baseline / Logistic Regression / Random Forest -> calibration -> model registry -> shadow mode. Sisipkan grafik perbandingan recall, F1, false alarm per jam, dan Brier score setelah eksperimen nyata selesai.",
        6,
    )
    add_caption(doc, "Gambar 6. Evidence gate pengembangan AI REKSA.")
    add_callout(doc, "Status AI saat proposal", "Service AI sudah terintegrasi dan mengembalikan provenance, tetapi belum ada klaim performa ML dari data lapangan. Data sintetis hanya digunakan untuk menguji pipeline, bukan sebagai bukti efektivitas keselamatan.", color=GOLD)
    add_page_break(doc)


def add_functional(doc):
    add_heading(doc, "6. ANALISIS FUNGSIONAL, CARA KERJA, DAN KINERJA", 1)
    add_heading(doc, "6.1 Kebutuhan Fungsional", 2)
    requirements = [
        "Helmet mengenali identitas hazard node yang valid melalui BLE advertising.",
        "Firmware menyaring RSSI serta membentuk state zona dengan hysteresis dan time persistence.",
        "Buzzer dan vibration motor aktif tanpa menunggu jaringan, backend, atau AI.",
        "Telemetry memuat message_id, timestamp, data gerakan, lingkungan, dan hazard terdekat.",
        "Record yang belum diakui disimpan, dikirim ulang, dan tidak diproses dua kali setelah ACK terlambat.",
        "Dashboard menampilkan status pekerja W01 dari hardware serta menandai data simulasi secara jelas.",
        "Candidate event dapat ditinjau dan diverifikasi supervisor sebelum menjadi verified near-miss.",
        "Layanan AI menampilkan mode model, versi, dan peran advisory-only.",
    ]
    for item in requirements:
        add_number(doc, item)
    add_heading(doc, "6.2 Cara Kerja Sistem", 2)
    flow = [
        ("Aktivasi", "Helmet melakukan inisialisasi sensor, BLE, Wi-Fi, sinkronisasi waktu, dan MQTT. Alarm lokal tidak menunggu seluruh koneksi selesai."),
        ("Sensing", "Helmet memindai hazard dan membaca MPU6050, DHT22, serta MQ135."),
        ("Klasifikasi", "RSSI diproses bersama history state untuk menentukan tingkat zona."),
        ("Peringatan", "Buzzer dan motor getar dijalankan sesuai level prioritas hazard."),
        ("Pencatatan", "Firmware membentuk payload dengan message_id unik dan memasukkannya ke buffer."),
        ("Pengiriman", "Payload dikirim melalui MQTT TLS/QoS 1. Tanpa ACK, record tetap berada di buffer."),
        ("Forwarding", "Bridge memasukkan record ke outbox, meneruskannya ke Pub/Sub, lalu memberi ACK setelah publish berhasil."),
        ("Pemrosesan cloud", "Backend menerima push, memperbarui state, mengarsipkan envelope, dan meminta advisory AI."),
        ("Tinjauan", "Supervisor membaca dashboard serta memverifikasi candidate event."),
    ]
    add_table(doc, ["Tahap", "Proses"], flow, [3.3, 13.5], font_size=9.2, first_col_bold=True)
    add_page_break(doc)
    add_heading(doc, "6.3 Taksonomi Kejadian Keselamatan", 2)
    add_table(
        doc,
        ["Tingkat", "Definisi Operasional", "Pihak yang Menetapkan"],
        [
            ("Telemetry", "Pembacaan periodik perangkat tanpa kesimpulan kejadian.", "Perangkat"),
            ("Unsafe proximity event", "Zona risiko terpenuhi setelah filtering dan persistence.", "Aturan lokal/backend"),
            ("Candidate near-miss", "Unsafe event dengan kombinasi durasi, hazard aktif, alarm, dan/atau gerakan mendadak.", "Aturan atau model advisory"),
            ("Verified near-miss", "Kandidat yang telah diperiksa dengan konteks operasional yang memadai.", "Supervisor K3"),
            ("False alarm / normal activity", "Kandidat yang tidak mewakili near-miss setelah verifikasi.", "Supervisor K3"),
        ],
        [3.5, 9.0, 4.3],
        font_size=8.9,
    )
    add_placeholder(
        doc,
        "PLACEHOLDER GAMBAR 7: TAKSONOMI DAN HUMAN-IN-THE-LOOP",
        "Buat diagram bertingkat Telemetry -> Unsafe Proximity Event -> Candidate Near-Miss -> Verified Near-Miss atau False Alarm. Tunjukkan bahwa supervisor adalah satu-satunya pihak yang menetapkan label final.",
        4,
    )
    add_caption(doc, "Gambar 7. Alur pembentukan dan verifikasi kejadian keselamatan.")
    add_heading(doc, "6.4 Batasan Kemampuan", 2)
    limits = [
        "BLE-RSSI digunakan untuk klasifikasi zona relatif, bukan pengukuran jarak presisi.",
        "REKSA bukan sistem antitabrakan tersertifikasi dan tidak menjamin kecelakaan dapat dicegah.",
        "Candidate near-miss bukan diagnosis otomatis; label final memerlukan verifikasi supervisor.",
        "Candidate fall/impact tidak membuktikan cedera atau jatuh secara definitif.",
        "DHT22 bukan alat WBGT, sedangkan MQ135 bukan alat AQI atau pengukur gas spesifik tanpa kalibrasi khusus.",
        "Dashboard cloud pada purwarupa masih memuat data seed untuk pekerja selain W01 dan menandainya sebagai simulation_data.",
        "Service AI saat ini masih baseline aturan dan belum mempunyai klaim performa lapangan.",
        "Service-account key gateway hanya untuk demonstrasi dan harus dirotasi setelah kegiatan.",
    ]
    for item in limits:
        add_bullet(doc, item)
    add_page_break(doc)
    add_heading(doc, "6.5 Kinerja dan Bukti Integrasi Saat Ini", 2)
    add_paragraph(
        doc,
        "Uji integrasi pada 17 Agustus 2026 dilakukan menggunakan helmet W01, hazard F01, HiveMQ Cloud, gateway lokal, dan layanan Google Cloud. Pengujian ini membuktikan keterhubungan komponen, bukan efektivitas keselamatan pada lingkungan industri penuh. Hasil yang telah diamati dirangkum sebagai berikut.",
    )
    add_table(
        doc,
        ["Objek Uji", "Hasil Teramati", "Interpretasi"],
        [
            ("MQTT TLS", "ESP32 menyelesaikan DNS, sinkronisasi NTP, dan terhubung ke broker port 8883.", "Transport aman berfungsi."),
            ("Alarm lokal", "Transisi SAFE ke CRITICAL memicu buzzer dan vibration motor saat MQTT belum tersambung.", "Prinsip offline-first terbukti secara fungsional."),
            ("Bridge", "Message_id disimpan sebagai Buffered dan memperoleh Pub/Sub message ID saat Forwarded.", "Gateway berhasil meneruskan telemetry."),
            ("Deduplication", "ID yang dikirim ulang tercatat sebagai Duplicate dan hanya di-ACK ulang.", "Satu ID tidak diteruskan dua kali setelah perbaikan."),
            ("Cloud ingest", "Cloud Run mencatat POST /api/v1/ingest/pubsub dengan HTTP 204.", "Push Pub/Sub diterima backend."),
            ("Dashboard", "W01 tampil dengan calculation_source HARDWARE_MQTT.", "Data perangkat nyata mencapai aplikasi."),
            ("Health endpoint", "Dashboard, backend, dan AI merespons HTTP 200.", "Deployment publik aktif."),
        ],
        [3.4, 8.1, 5.3],
        font_size=8.7,
    )
    add_heading(doc, "6.6 Rencana Pengukuran Performansi", 2)
    add_table(
        doc,
        ["Metrik", "Definisi", "Metode", "Kriteria Promosi"],
        [
            ("P95 warning latency", "Waktu dari kondisi zona terpenuhi hingga aktuator aktif", "Timestamp firmware pada pengulangan terkontrol", "Dilaporkan sebagai hasil, bukan target semata"),
            ("Missed warning rate", "Proporsi skenario bahaya tanpa alarm", "Skenario berlabel dan video ground truth", "Serendah mungkin; interval kepercayaan disertakan"),
            ("False alarm per jam", "Alarm pada skenario aman per waktu operasi", "Sesi negatif berdurasi tetap", "Model tidak boleh lebih buruk dari baseline"),
            ("Recall dan F1", "Kemampuan menemukan kandidat positif dan keseimbangannya", "Held-out participant/session", "Model harus mengungguli baseline"),
            ("Brier score", "Kualitas probabilitas", "Prediksi pada holdout", "Lebih rendah dari kandidat lain"),
            ("Data delivery", "Generated, buffered, resent, ACK, duplicate, dropped", "Serial dan bridge log", "Klaim zero loss hanya jika dropped=0"),
            ("Daya", "Arus idle, rata-rata, puncak, dan durasi", "USB power meter dan uji baterai", "Sesuai durasi demo dan batas termal"),
        ],
        [3.1, 4.8, 5.1, 3.8],
        font_size=8.1,
    )
    add_page_break(doc)


def add_plan(doc):
    add_heading(doc, "7. RENCANA IMPLEMENTASI DAN PERKEMBANGAN PENGERJAAN", 1)
    add_heading(doc, "7.1 Status Pengerjaan Saat Proposal", 2)
    add_paragraph(
        doc,
        "Purwarupa telah melampaui integrasi dasar 50 persen. Helmet dan hazard node berfungsi, alarm lokal telah diuji saat MQTT belum tersedia, koneksi HiveMQ TLS berjalan, dan pipeline cloud telah dipublikasikan. Pekerjaan utama yang belum selesai bukan lagi menyambungkan komponen, tetapi menghasilkan bukti kuantitatif yang cukup untuk kalibrasi zona dan promosi model AI.",
    )
    add_table(
        doc,
        ["Komponen", "Status", "Bukti / Catatan"],
        [
            ("Hazard BLE advertiser", "Selesai", "F01 dikenali oleh helmet."),
            ("Helmet scanner dan alarm", "Selesai", "State SAFE sampai CRITICAL dan alarm multimodal."),
            ("MPU6050, DHT22, MQ135", "Terintegrasi", "Telemetry perangkat nyata masuk dashboard."),
            ("HiveMQ TLS + QoS 1", "Selesai", "Koneksi port 8883 dan test publish/subscribe berhasil."),
            ("Buffer, ACK, retry", "Selesai", "Firmware menahan record sampai ACK."),
            ("Bridge durable + dedup", "Selesai", "Duplicate di-ACK tanpa forward kedua."),
            ("Pub/Sub, Cloud Run, Firestore", "Selesai", "POST ingest 204 dan health 200."),
            ("Firebase dashboard", "Selesai untuk demo", "W01 memakai HARDWARE_MQTT; data lain masih seed."),
            ("Dataset perangkat nyata", "Berjalan", "Perlu sesi 5-10 Hz dan label per trajectory."),
            ("Model ML terpilih", "Belum", "Menunggu evaluasi baseline, LR, dan RF."),
            ("Kalibrasi multi-lokasi", "Belum", "Perlu orientasi, rak logam, dan hazard bergerak."),
        ],
        [5.2, 3.0, 8.6],
        font_size=8.6,
    )
    add_heading(doc, "7.2 Tahapan Penyelesaian", 2)
    add_table(
        doc,
        ["Tahap", "Fokus", "Luaran"],
        [
            ("Tahap A: Stabilitas perangkat", "Kalibrasi RSSI, daya, aktuator, dan ergonomi", "Parameter state machine dan log latency"),
            ("Tahap B: Evidence IoT", "Disconnect, reconnect, outbox, ACK, dan dropped record", "Tabel delivery dan video bukti"),
            ("Tahap C: Dataset", "Skenario positif/negatif pada beberapa peserta dan sesi", "Raw log, data dictionary, dan label"),
            ("Tahap D: Model", "Baseline, Logistic Regression, Random Forest, calibration", "Model evidence dan artifact terpilih"),
            ("Tahap E: Shadow mode", "Model memberi advisory tanpa mengendalikan alarm", "Perbandingan rekomendasi dengan supervisor"),
            ("Tahap F: Finalisasi", "Casing, dashboard, performansi, video, dan laporan", "Purwarupa 100 persen dan paket demo"),
        ],
        [3.9, 7.2, 5.7],
        font_size=8.8,
    )
    add_page_break(doc)
    add_heading(doc, "7.3 Timeline", 2)
    add_table(
        doc,
        ["Minggu", "Kegiatan", "Indikator Selesai"],
        [
            ("1-2", "Kalibrasi zona, uji orientasi, dan pengukuran warning latency", "Grafik RSSI mentah/terfilter dan tabel P95"),
            ("3", "Uji disconnect, reconnect, retransmisi, dan deduplikasi", "Generated = ACK + buffered + dropped yang dapat dijelaskan"),
            ("4-5", "Perekaman dataset 5-10 Hz pada skenario terkontrol", "Dataset lintas peserta, sesi, dan trajectory"),
            ("6", "Pelatihan dan evaluasi tiga kandidat model", "Model evidence, confusion matrix, calibration"),
            ("7", "Deploy pemenang dalam shadow mode", "Provenance model tampil di dashboard"),
            ("8", "Uji daya, casing, kenyamanan, dan stabilitas berjam-jam", "Tabel konsumsi daya dan catatan kegagalan"),
            ("9", "Penyempurnaan dashboard dan alur verifikasi", "Demo end-to-end tanpa data hard-coded"),
            ("10", "Video, laporan akhir, dan rehearsal tanya-jawab", "Paket final dan daftar bukti klaim"),
        ],
        [2.2, 8.4, 6.2],
        font_size=8.7,
    )
    add_heading(doc, "7.4 Estimasi Biaya Purwarupa", 2)
    add_table(
        doc,
        ["Kelompok", "Komponen", "Kisaran Biaya", "Catatan"],
        [
            ("Helmet", "ESP32, MPU6050, DHT22, MQ135, aktuator, driver", "Rp350.000-Rp550.000", "Tidak termasuk helm bersertifikasi"),
            ("Hazard node", "ESP32, catu daya, casing", "Rp150.000-Rp250.000", "Satu node demonstrasi"),
            ("Mekanik", "Casing 3D, kabel, konektor, mounting", "Rp150.000-Rp300.000", "Bergantung iterasi desain"),
            ("Cloud", "HiveMQ Serverless dan layanan Google Cloud", "Free tier / usage-based", "Budget alert Rp20.000 per bulan untuk demo"),
            ("Pengujian", "Power meter, material rig, dokumentasi", "Rp150.000-Rp300.000", "Dapat memakai alat laboratorium"),
        ],
        [3.0, 6.0, 3.4, 4.4],
        font_size=8.6,
    )
    add_paragraph(doc, "Angka merupakan estimasi awal dan harus diganti dengan nota pembelian aktual sebelum laporan akhir.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5)
    add_page_break(doc)
    add_heading(doc, "7.5 Risiko Teknis dan Mitigasi", 2)
    add_table(
        doc,
        ["Risiko", "Dampak", "Mitigasi", "Bukti Verifikasi"],
        [
            ("RSSI fluktuatif", "State tidak stabil", "Filtering, hysteresis, persistence, kalibrasi per lokasi", "Grafik dan jumlah transisi salah"),
            ("Tubuh/rak menghalangi BLE", "Missed warning", "Uji orientasi dan penempatan node", "Recall per orientasi"),
            ("TLS kekurangan heap", "MQTT gagal", "NimBLE, optimasi memori, NTP sebelum handshake", "Free heap dan TLS log"),
            ("Wi-Fi putus", "Telemetry tertunda", "Buffer, retry, outbox, ACK", "Log reconnect dan dropped"),
            ("ACK hilang", "Retransmisi", "Persistent delivered-ID dan re-ACK duplicate", "Satu forward per ID"),
            ("Model overfit", "Advisory menyesatkan", "Grouped holdout dan baseline comparison", "Model evidence"),
            ("Cloud tidak tersedia", "Dashboard tidak real-time", "Alarm lokal tetap aktif", "Demo cabut internet"),
            ("Key gateway bocor", "Akses cloud tidak sah", "Git ignore, least privilege, rotasi setelah lomba", "Audit IAM"),
            ("Daya turun", "Reset atau sensor salah", "Driver, regulator, kapasitor, uji arus puncak", "Log brownout dan power meter"),
        ],
        [3.0, 3.3, 6.2, 4.3],
        font_size=8.0,
    )
    add_heading(doc, "7.6 Strategi Validasi Klaim", 2)
    add_paragraph(
        doc,
        "Setiap klaim pada laporan akhir akan dipasangkan dengan satu artefak bukti, misalnya log serial, log bridge, tangkapan Cloud Run, rekaman video, tabel eksperimen, atau artifact model. Tim tidak akan menggunakan frasa real-time, akurat, zero data loss, maupun prediktif tanpa definisi pengukuran. Strategi ini penting karena kualitas karya bukan hanya ditentukan oleh fitur yang berhasil pada satu demonstrasi, tetapi juga oleh kemampuan menjelaskan batas, kegagalan, dan proses perbaikannya.",
    )
    add_page_break(doc)


def add_photos(doc):
    add_heading(doc, "8. FOTO DAN PENJELASAN HASIL IMPLEMENTASI PURWARUPA", 1)
    add_paragraph(
        doc,
        "Bagian ini memuat bukti yang sudah tersedia serta placeholder untuk dokumentasi yang perlu diambil ulang menjelang pengumpulan. Setiap foto final sebaiknya mempunyai pencahayaan yang cukup, latar sederhana, resolusi tinggi, dan label komponen yang tetap terbaca setelah dokumen diekspor ke PDF.",
    )
    photos = [MEDIA / "image1.jpg", MEDIA / "image2.jpg"]
    for idx, photo in enumerate(photos, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        add_picture_with_alt(
            p.add_run(),
            photo,
            Cm(11.6),
            f"Purwarupa Smart Helmet REKSA dari sudut {'samping' if idx == 1 else 'depan'}",
        )
        add_caption(doc, f"Gambar {7+idx}. Purwarupa Smart Helmet REKSA dari sudut {'samping' if idx == 1 else 'depan'}.")
    add_callout(
        doc,
        "Pemeriksaan foto",
        "Pastikan foto yang dipakai benar-benar menampilkan purwarupa tim. Jika casing atau posisi komponen berubah setelah dokumen ini dibuat, ganti kedua foto di atas agar konsisten dengan perangkat yang didemonstrasikan.",
        color=GOLD,
    )
    add_page_break(doc)
    add_heading(doc, "8.1 Dokumentasi Perangkat dan Pengujian yang Harus Ditambahkan", 2)
    placeholders = [
        ("PLACEHOLDER FOTO A: HELMET TERPAKAI", "Foto anggota tim memakai helm dari sisi depan dan samping. Tampilkan posisi casing, jalur kabel yang aman, serta tidak adanya bagian tajam atau terbuka."),
        ("PLACEHOLDER FOTO B: RANGKAIAN INTERNAL", "Close-up ESP32, MPU6050, DHT22, MQ135, driver buzzer, driver motor, konektor, dan sumber daya. Beri callout nomor pada setiap komponen."),
        ("PLACEHOLDER FOTO C: ACTIVE HAZARD NODE F01", "Foto ESP32 advertiser terpasang pada miniatur forklift atau objek hazard. Nama REKSA_HAZARD_F01 harus terlihat pada label fisik."),
        ("PLACEHOLDER FOTO D: SETUP UJI ZONA", "Foto jarak dan orientasi uji menggunakan penanda lantai. Tampilkan posisi SAFE, MODERATE, HIGH, dan CRITICAL berdasarkan hasil kalibrasi."),
    ]
    for i, (title, instruction) in enumerate(placeholders, start=10):
        add_placeholder(doc, title, instruction, 3)
        add_caption(doc, f"Gambar {i}. Dokumentasi yang wajib diganti dengan foto aktual sebelum pengumpulan.")
    add_page_break(doc)
    add_heading(doc, "8.2 Dokumentasi Perangkat Lunak dan Cloud yang Harus Ditambahkan", 2)
    placeholders = [
        ("PLACEHOLDER BUKTI E: SERIAL MONITOR OFFLINE-FIRST", "Tampilkan METRIC ZONE_TRANSITION dan ALARM_APPLIED ketika mqtt=0. Sertakan timestamp dan jangan memotong baris penting."),
        ("PLACEHOLDER BUKTI F: HIVEMQ DAN BRIDGE", "Gabungkan status cluster Running dengan log Connected, Buffered, Forwarded, dan Duplicate re-acknowledged. Jangan tampilkan password."),
        ("PLACEHOLDER BUKTI G: CLOUD RUN DAN PUB/SUB", "Tampilkan status healthy, POST ingest 204, topic, dan subscription. Samarkan token serta informasi rahasia."),
        ("PLACEHOLDER BUKTI H: DASHBOARD FIREBASE", "Tampilkan halaman Live Monitoring saat W01 berubah mengikuti hazard F01. Pastikan calculation_source HARDWARE_MQTT dan label simulation data terlihat jujur."),
        ("PLACEHOLDER GRAFIK I: HASIL UJI", "Sisipkan grafik RSSI mentah vs terfilter, latency distribution, false alarm per jam, dan tabel delivery. Gunakan data aktual serta cantumkan jumlah sesi."),
        ("PLACEHOLDER GRAFIK J: MODEL EVIDENCE", "Setelah dataset nyata tersedia, tampilkan recall, F1, Brier score, confusion matrix, serta calibration curve untuk baseline, Logistic Regression, dan Random Forest."),
    ]
    for i, (title, instruction) in enumerate(placeholders, start=14):
        add_placeholder(doc, title, instruction, 3)
        add_caption(doc, f"Gambar {i}. Placeholder bukti teknis yang harus diganti sebelum proposal final dikirim.")
    add_page_break(doc)


def add_video_and_refs(doc):
    add_heading(doc, "9. TAUTAN VIDEO PROSES PENGEMBANGAN MODEL KARYA INOVASI", 1)
    add_paragraph(
        doc,
        "Video menjelaskan masalah, arsitektur, dan bukti bahwa ketiga elemen kategori telah diterapkan. Urutan yang disarankan adalah: pengenalan masalah oleh narator, demonstrasi hazard mendekat, alarm lokal, tampilan message_id, forwarding cloud, perubahan dashboard, lalu pemutusan internet untuk membuktikan alarm tetap bekerja. Video juga perlu menyatakan bahwa AI masih advisory-only dan model ML hanya digunakan setelah lolos evaluasi terhadap baseline.",
    )
    add_table(
        doc,
        ["Segmen", "Durasi", "Isi"],
        [
            ("Pembuka", "0:00-0:20", "Masalah dan manfaat bagi pekerja serta supervisor"),
            ("Perangkat", "0:20-0:55", "Helmet, hazard node, sensor, dan alarm lokal"),
            ("IoT-cloud", "0:55-1:35", "HiveMQ, bridge, Pub/Sub, Cloud Run, dashboard"),
            ("AI", "1:35-2:05", "Baseline, fitur temporal, evidence gate, human-in-the-loop"),
            ("Failure demo", "2:05-2:35", "Cabut internet, alarm tetap aktif, reconnect dan retransmisi"),
            ("Penutup", "2:35-3:00", "Dampak, roadmap, dan identitas tim"),
        ],
        [3.2, 2.5, 11.1],
        font_size=9.1,
    )
    add_callout(
        doc,
        "Tautan video saat ini",
        "https://youtu.be/aq8y6m_AiCI\n\nGanti tautan ini apabila tim mengunggah versi final baru. Judul YouTube harus mengikuti format resmi GEMASTIK XIX dan durasi tidak melebihi tiga menit.",
        color=GOLD,
    )
    add_page_break(doc)
    add_heading(doc, "10. DAFTAR PUSTAKA", 1)
    refs = [
        "BPJS Ketenagakerjaan. (2025). Dukung Pemerintah Perkuat Budaya K3, BPJS Ketenagakerjaan Lakukan Berbagai Upaya Promotif Preventif. https://www.bpjsketenagakerjaan.go.id/berita/29541/",
        "Espressif Systems. (2023). ESP32 Series Datasheet: Wi-Fi & Bluetooth SoC. Espressif Systems.",
        "Google Cloud. (2026). Authentication for push subscriptions. Google Cloud Documentation. https://cloud.google.com/pubsub/docs/authenticate-push-subscriptions",
        "Google Cloud. (2026). Use Pub/Sub with Cloud Run. Google Cloud Documentation. https://cloud.google.com/run/docs/tutorials/pubsub",
        "Inagaki, M., Nagata, T., Odagami, K., Adi, N. P., dan Mori, K. (2024). Relationship between a company's adequate response to near-misses and occupational accidents: A 1-year prospective cohort study. Journal of Occupational Health, 66(1), uiae053. https://doi.org/10.1093/joccuh/uiae053",
        "International Labour Organization. (2024). Indonesia launches its five-year National Occupational Safety and Health Programme 2024-2029. ILO.",
        "International Organization for Standardization. (2018). ISO 45001:2018 Occupational health and safety management systems: Requirements with guidance for use. ISO.",
        "Kim, Y., Baek, J., dan Choi, Y. (2021). Smart helmet-based personnel proximity warning system for improving underground mine safety. Applied Sciences, 11(10), 4342. https://doi.org/10.3390/app11104342",
        "Kyung, M., Lee, S.-J., Dancu, C., dan Hong, O. (2023). Underreporting of workers' injuries or illnesses and contributing factors: A systematic review. BMC Public Health, 23, 558. https://doi.org/10.1186/s12889-023-15487-0",
        "Lee, P., Kim, H., Zitouni, M. S., Khandoker, A., Jelinek, H. F., Hadjileontiadis, L., Lee, U., dan Jeong, Y. (2022). Trends in smart helmets with multimodal sensing for health and safety: Scoping review. JMIR mHealth and uHealth, 10(11), e40797. https://doi.org/10.2196/40797",
        "OASIS. (2019). MQTT Version 5.0. OASIS Standard. https://docs.oasis-open.org/mqtt/mqtt/v5.0/mqtt-v5.0.html",
        "Pusat Prestasi Nasional. (2026). Pedoman Pagelaran Mahasiswa Nasional Bidang Teknologi Informasi dan Komunikasi GEMASTIK XIX Tahun 2026. Kementerian Pendidikan Tinggi, Sains, dan Teknologi.",
        "Vaccari, L., Coruzzolo, A. M., Lolli, F., dan Sellitto, M. A. (2024). Indoor positioning systems in logistics: A review. Logistics, 8(4), 126. https://doi.org/10.3390/logistics8040126",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(ref)
        set_run_font(r, size=9.7)


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "REKSA - Proposal GEMASTIK XIX 2026"
    props.subject = "Piranti Cerdas, Sistem Benam & IoT"
    props.author = "Tim Restu Bundo, Universitas Gunadarma"
    props.keywords = "REKSA, GEMASTIK, AI-IoT, smart helmet, K3, near-miss"
    props.comments = "Naskah revisi berbasis implementasi aktual. Placeholder visual harus diganti sebelum unggah."


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    set_update_fields(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_toc(doc)
    add_abstract(doc)
    add_background(doc)
    add_urgency(doc)
    add_method(doc)
    add_design(doc)
    add_functional(doc)
    add_plan(doc)
    add_photos(doc)
    add_video_and_refs(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
